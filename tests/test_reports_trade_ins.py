"""
Trade-ins were completely invisible outside their own page — not in
Reports, not in the Month-End Workbook, and (a real, connected bug found
while building this) not accounted for in Cash Up's drawer reconciliation
either, even though a cash trade-in payout is real money leaving the
till. These tests cover all three.
"""
import io
from datetime import datetime, timedelta

from openpyxl import load_workbook

from app.models import TradeIn, Customer
from app.main import todays_cash_card_totals


def _accept_trade_in(client, device, amount, payout_method="cash", **extra):
    data = {
        "customer_id": "", "new_customer_name": "", "new_customer_phone": "",
        "device": device, "imei": "", "condition": "",
        "offered_amount": str(amount), "payout_method": payout_method, "notes": "",
    }
    data.update(extra)
    return client.post("/trade-ins/add", data=data, follow_redirects=False)


# ── Reports page ─────────────────────────────────────────────────────

def test_reports_shows_trade_in_summary_this_month(owner_client, db_session):
    _accept_trade_in(owner_client, "Reports Cash Device", "40.00", "cash")
    resp = owner_client.get("/reports")
    assert resp.status_code == 200
    assert "Trade-Ins" in resp.text
    assert "Cash Paid Out" in resp.text
    assert "Store Credit Issued" in resp.text


def test_reports_trade_in_totals_separate_cash_from_store_credit(owner_client, db_session):
    _accept_trade_in(owner_client, "Reports Cash Split Device", "40.00", "cash")
    _accept_trade_in(owner_client, "Reports Credit Split Device", "60.00", "store_credit")

    now = datetime.utcnow()
    month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    this_month = [t for t in db_session.query(TradeIn).all() if t.created_at >= month_start]
    cash_paid = round(sum(t.offered_amount for t in this_month if t.payout_method == "cash"), 2)
    credit_issued = round(sum(t.offered_amount for t in this_month if t.payout_method == "store_credit"), 2)

    resp = owner_client.get("/reports")
    assert f"${cash_paid:.2f}" in resp.text
    assert f"${credit_issued:.2f}" in resp.text


# ── Month-End Workbook ──────────────────────────────────────────────

def test_month_end_workbook_has_a_trade_ins_sheet(owner_client, db_session):
    _accept_trade_in(owner_client, "Workbook Sheet Device", "75.00", "cash")
    resp = owner_client.get("/export/xlsx/month-end")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Trade-Ins" in wb.sheetnames


def test_month_end_workbook_trade_ins_sheet_lists_the_device_and_amount(owner_client, db_session):
    _accept_trade_in(owner_client, "Workbook Row Device", "82.50", "cash")
    resp = owner_client.get("/export/xlsx/month-end")
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Trade-Ins"]
    rows = list(ws.iter_rows(values_only=True))
    devices = [r[3] for r in rows[1:] if r[3]]
    assert "Workbook Row Device" in devices


def test_month_end_workbook_trade_ins_sheet_has_totals_footer(owner_client, db_session):
    _accept_trade_in(owner_client, "Workbook Cash Footer Device", "20.00", "cash")
    _accept_trade_in(owner_client, "Workbook Credit Footer Device", "30.00", "store_credit")
    resp = owner_client.get("/export/xlsx/month-end")
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Trade-Ins"]
    labels = [row[0] for row in ws.iter_rows(values_only=True) if row[0]]
    assert "Cash Paid Out" in labels
    assert "Store Credit Issued" in labels


def test_month_end_workbook_trade_ins_excludes_other_months(owner_client, db_session):
    old = TradeIn(
        number="TR-OLDMONTH", device="Last Year Device", offered_amount=99.00,
        payout_method="cash", staff_name="Owner",
        created_at=datetime.utcnow() - timedelta(days=400),
    )
    db_session.add(old)
    db_session.commit()

    resp = owner_client.get("/export/xlsx/month-end")
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Trade-Ins"]
    devices = [row[3] for row in ws.iter_rows(values_only=True) if row[3]]
    assert "Last Year Device" not in devices


# ── Cash Up — the connected fix ─────────────────────────────────────

def test_cash_trade_in_payout_reduces_expected_cash(owner_client, db_session):
    today_str = datetime.utcnow().strftime("%Y-%m-%d")
    before = todays_cash_card_totals(db_session, today_str)["cash_sales"]

    _accept_trade_in(owner_client, "Cashup Cash Payout Device", "50.00", "cash")

    after = todays_cash_card_totals(db_session, today_str)["cash_sales"]
    assert round(before - after, 2) == 50.00


def test_store_credit_trade_in_does_not_affect_cash_up(owner_client, db_session):
    today_str = datetime.utcnow().strftime("%Y-%m-%d")
    before = todays_cash_card_totals(db_session, today_str)

    _accept_trade_in(owner_client, "Cashup Store Credit Device", "50.00", "store_credit")

    after = todays_cash_card_totals(db_session, today_str)
    assert after["cash_sales"] == before["cash_sales"]
    assert after["card_sales"] == before["card_sales"]


def test_cash_trade_in_from_a_different_day_does_not_affect_today(db_session):
    today_str = datetime.utcnow().strftime("%Y-%m-%d")
    before = todays_cash_card_totals(db_session, today_str)["cash_sales"]

    old = TradeIn(
        number="TR-YESTERDAY", device="Yesterday Device", offered_amount=999.00,
        payout_method="cash", staff_name="Owner",
        created_at=datetime.utcnow() - timedelta(days=1),
    )
    db_session.add(old)
    db_session.commit()

    after = todays_cash_card_totals(db_session, today_str)["cash_sales"]
    assert after == before
